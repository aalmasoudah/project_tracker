"""Phase 13 report permissions, scoping, formats, and privacy tests."""

from datetime import date
from importlib import import_module
from io import BytesIO
from typing import cast

import pytest
from django.contrib.auth.models import Group
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from apps.accounts.models import User
from apps.attendance.models import AttendanceEntry, AttendanceReview
from apps.attendance.services import (
    issue_trainer_link,
    review_attendance,
    submit_attendance,
)
from apps.organizations.models import Department
from apps.projects.models import Project
from tests.integration.test_phase9_attendance import (
    create_one_session,
    setup_course,
)

PHASE13_ROLE_PERMISSIONS = cast(
    "dict[str, set[str]]",
    import_module(
        "apps.reports.migrations.0002_seed_phase13_permissions"
    ).ROLE_PERMISSIONS,
)


def role_user(username: str, role: str, department: Department) -> User:
    user = User.objects.create_user(
        username=username,
        email=f"{username}@example.test",
        display_name=username.replace("-", " ").title(),
        department=department,
        password="fictional-phase13-password-9321",
    )
    user.groups.add(Group.objects.get(name=role))
    return user


def create_project(code: str, manager: User, department: Department) -> Project:
    return Project.objects.create(
        code=code,
        name_ar=f"مشروع {code}",
        name_en=f"Project {code}",
        department=department,
        manager=manager,
        status=Project.Status.ACTIVE,
        priority=Project.Priority.HIGH,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
        created_by=manager,
        updated_by=manager,
    )


def workbook_values(payload: bytes) -> set[str]:
    workbook = load_workbook(BytesIO(payload), read_only=True)
    sheet = cast(Worksheet, workbook.active)
    return {
        str(cell)
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    }


@pytest.mark.integration
@pytest.mark.django_db
def test_seeded_report_permissions_exactly_match_approved_matrix() -> None:
    for role, expected in PHASE13_ROLE_PERMISSIONS.items():
        group = Group.objects.get(name=role)
        actual = set(
            group.permissions.filter(content_type__app_label="reports").values_list(
                "codename", flat=True
            )
        )
        assert actual == expected


@pytest.mark.integration
@pytest.mark.security
@pytest.mark.django_db
def test_report_index_and_forged_formats_follow_role_policy() -> None:
    department = Department.objects.create(
        code="REP13", name_ar="إدارة التقارير", name_en="Reports"
    )
    supervisor = role_user("report-supervisor", "supervisor", department)
    employee = role_user("report-employee", "employee", department)
    manager = role_user("report-manager", "project_manager", department)
    client = Client()

    client.force_login(employee)
    assert client.get(reverse("reports:index")).status_code == 403
    client.force_login(supervisor)
    response = client.get(reverse("reports:index"))
    choices = {
        code for code, _label in response.context["form"].fields["report_type"].choices
    }
    assert response.status_code == 200
    assert choices == {"course_attendance", "project_attendance"}

    client.force_login(manager)
    forged = client.post(
        reverse("reports:generate"),
        {
            "report_type": "overdue_tasks",
            "output_format": "csv",
            "locale": "en",
            "date_from": "2026-01-01",
            "date_to": "2026-01-31",
        },
    )
    assert forged.status_code == 400
    excessive_range = client.post(
        reverse("reports:generate"),
        {
            "report_type": "project_progress",
            "output_format": "xlsx",
            "locale": "en",
            "date_from": "2025-01-01",
            "date_to": "2026-12-31",
        },
    )
    assert excessive_range.status_code == 400


@pytest.mark.integration
@pytest.mark.security
@pytest.mark.django_db
def test_project_progress_export_is_scoped_and_not_retained() -> None:
    department = Department.objects.create(
        code="SCP13", name_ar="إدارة النطاق", name_en="Scope"
    )
    manager = role_user("scoped-report-manager", "project_manager", department)
    outsider = role_user("hidden-report-manager", "project_manager", department)
    visible = create_project("REP-VISIBLE", manager, department)
    hidden = create_project("REP-HIDDEN", outsider, department)
    client = Client()
    client.force_login(manager)

    response = client.post(
        reverse("reports:generate"),
        {
            "report_type": "project_progress",
            "output_format": "xlsx",
            "locale": "en",
            "date_from": "2026-01-01",
            "date_to": "2026-12-31",
        },
    )

    assert response.status_code == 200
    assert response["Cache-Control"] == "no-store, private"
    assert response["Content-Disposition"].endswith('"project-progress-en.xlsx"')
    values = workbook_values(response.content)
    assert visible.code in values
    assert hidden.code not in values

    forged = client.post(
        reverse("reports:generate"),
        {
            "report_type": "project_progress",
            "output_format": "xlsx",
            "project": hidden.pk,
            "locale": "en",
        },
    )
    assert forged.status_code == 400


@pytest.mark.integration
@pytest.mark.security
@pytest.mark.django_db
def test_attendance_export_is_aggregate_only_and_pdf_is_available() -> None:
    manager, supervisor, course, trainer = setup_course()
    session = create_one_session(manager, course, trainer)
    link, _token = issue_trainer_link(
        actor=manager,
        session=session,
        lifetime_hours=1,
    )
    entries = {
        participant.pk: (str(AttendanceEntry.Value.PRESENT), "")
        for participant in session.participants.all()
    }
    submission = submit_attendance(
        link=link,
        entries=entries,
        trainer_notes="Fictional trainer note",
        evidence_files=[],
    )
    review_attendance(
        actor=supervisor,
        submission=submission,
        decision=AttendanceReview.Decision.APPROVED,
    )
    client = Client()
    client.force_login(supervisor)
    data = {
        "report_type": "course_attendance",
        "course": course.pk,
        "locale": "ar",
        "date_from": "2026-08-01",
        "date_to": "2026-08-31",
    }
    excel = client.post(
        reverse("reports:generate"),
        {**data, "output_format": "xlsx"},
    )
    pdf = client.post(
        reverse("reports:generate"),
        {**data, "output_format": "pdf"},
    )

    assert excel.status_code == 200
    values = workbook_values(excel.content)
    assert "2" in values
    assert "متدرب خيالي 0" not in values
    assert "متدرب خيالي 1" not in values
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")
    assert pdf["Content-Type"] == "application/pdf"


@pytest.mark.integration
@pytest.mark.django_db
def test_every_approved_report_format_generates_in_both_languages() -> None:
    manager, _supervisor, course, _trainer = setup_course()
    client = Client()
    client.force_login(manager)
    approved: tuple[tuple[str, str, dict[str, int]], ...] = (
        ("project_progress", "pdf", {}),
        ("project_progress", "xlsx", {}),
        ("overdue_tasks", "pdf", {}),
        ("overdue_tasks", "xlsx", {}),
        ("course_attendance", "pdf", {"course": course.pk}),
        ("course_attendance", "xlsx", {"course": course.pk}),
        ("project_attendance", "pdf", {"project": course.project_id}),
        ("project_attendance", "xlsx", {"project": course.project_id}),
    )

    for language_code in ("ar", "en"):
        for report_type, output_format, scope in approved:
            response = client.post(
                reverse("reports:generate"),
                {
                    "report_type": report_type,
                    "output_format": output_format,
                    "locale": language_code,
                    "date_from": "2026-08-01",
                    "date_to": "2026-08-31",
                    **scope,
                },
            )

            assert response.status_code == 200
            assert response["Content-Disposition"].endswith(
                f'-{language_code}.{output_format}"'
            )
            expected_magic = b"%PDF-" if output_format == "pdf" else b"PK"
            assert response.content.startswith(expected_magic)
