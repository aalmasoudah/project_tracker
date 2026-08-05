"""Pure Phase 13 date, spreadsheet-safety, and renderer tests."""

import re
from datetime import date
from io import BytesIO
from typing import Any, cast

import pytest
from django.db.models import QuerySet
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from apps.reports.datasets import ReportDocument, ReportTooLargeError, _bounded
from apps.reports.dates import format_dual_date
from apps.reports.renderers import (
    brand_name,
    fit_dimensions,
    render_pdf,
    render_xlsx,
    safe_spreadsheet_value,
)


def sample_document(*, formula_value: str = "PRJ-13") -> ReportDocument:
    return ReportDocument(
        title="تقرير تقدم المشاريع",
        subtitle="جميع المشاريع المصرح بها",
        headers=("الرمز", "المشروع", "التاريخ"),
        rows=((formula_value, "مشروع التحول", "2026-07-30 م / 1448-02-15 هـ"),),
        filename_stem="project-progress",
        sheet_name="تقدم المشاريع",
    )


@pytest.mark.unit
def test_dual_date_uses_western_digits_and_hijri_markers() -> None:
    arabic = format_dual_date(date(2026, 7, 30), "ar")
    english = format_dual_date(date(2026, 7, 30), "en")

    assert arabic.startswith("2026-07-30 م / ")
    assert arabic.endswith(" هـ")
    assert english.startswith("2026-07-30 G / ")
    assert english.endswith(" AH")
    assert re.fullmatch(r"[0-9\- م/هـ]+", arabic)
    assert not re.search(r"[\u0660-\u0669]", arabic)


@pytest.mark.unit
@pytest.mark.parametrize("prefix", ("=", "+", "-", "@"))
def test_spreadsheet_formula_prefixes_are_neutralized(prefix: str) -> None:
    assert safe_spreadsheet_value(f"{prefix}SUM(A1:A2)").startswith("'")
    assert safe_spreadsheet_value("ordinary text") == "ordinary text"


@pytest.mark.unit
def test_arabic_excel_is_rtl_branded_and_formula_safe() -> None:
    payload = render_xlsx(sample_document(formula_value="=2+2"), language_code="ar")
    workbook = load_workbook(BytesIO(payload))
    sheet = cast(Worksheet, workbook.active)

    assert sheet.sheet_view.rightToLeft
    assert sheet["A6"].value == "'=2+2"
    assert sheet["A5"].fill.fgColor.rgb == "000A400C"
    images = cast(Any, sheet)._images
    assert len(images) == 1
    assert images[0].width / images[0].height == pytest.approx(2472 / 2649)
    assert sheet.oddFooter is not None
    assert sheet.oddFooter.center is not None
    assert sheet.oddFooter.center.text == "إنسايت بروجكتس"
    assert workbook.properties.creator == "إنسايت بروجكتس"


@pytest.mark.unit
def test_company_name_and_logo_dimensions_are_exact_and_proportional() -> None:
    assert brand_name("en") == "Insight Projects"
    assert brand_name("ar") == "إنسايت بروجكتس"
    assert brand_name("ar-sa") == "إنسايت بروجكتس"

    width, height = fit_dimensions(2472, 2649, max_width=170, max_height=95)
    assert width / height == pytest.approx(2472 / 2649)
    assert width <= 170
    assert height <= 95


@pytest.mark.unit
@pytest.mark.parametrize(
    ("language_code", "expected_message"),
    (("ar", "لا توجد بيانات مطابقة."), ("en", "No matching data.")),
)
def test_empty_excel_has_a_clear_localized_message(
    language_code: str,
    expected_message: str,
) -> None:
    document = sample_document()
    empty_document = ReportDocument(
        title=document.title,
        subtitle=document.subtitle,
        headers=document.headers,
        rows=(),
        filename_stem=document.filename_stem,
        sheet_name=document.sheet_name,
    )

    payload = render_xlsx(empty_document, language_code=language_code)
    workbook = load_workbook(BytesIO(payload))
    sheet = cast(Worksheet, workbook.active)

    assert sheet["A6"].value == expected_message
    assert "A6:C6" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}


@pytest.mark.unit
def test_arabic_pdf_embeds_report_content_without_temporary_files() -> None:
    payload = render_pdf(sample_document(), language_code="ar")

    assert payload.startswith(b"%PDF-")
    assert len(payload) > 10_000


@pytest.mark.unit
def test_report_dataset_rejects_more_than_five_thousand_rows() -> None:
    oversized = cast(QuerySet[Any], list(range(5_001)))
    with pytest.raises(ReportTooLargeError):
        _bounded(oversized)
