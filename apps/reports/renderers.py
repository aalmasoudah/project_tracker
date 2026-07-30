"""In-memory branded PDF and Excel report rendering."""

import re
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any, Final, Literal, cast

# These two focused text-layout libraries do not publish PEP 561 metadata.
import arabic_reshaper  # type: ignore[import-untyped]
from bidi.algorithm import get_display  # type: ignore[import-untyped]
from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Flowable,
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from apps.reports.datasets import ReportDocument

PRIMARY: Final = "0A400C"
PRIMARY_DARK: Final = "063008"
SAGE: Final = "819067"
STONE: Final = "B1AB86"
SOFT_STONE: Final = "F3F1E8"
TEXT: Final = "243127"
FORMULA_PREFIXES: Final = ("=", "+", "-", "@")
ARABIC_PATTERN: Final = re.compile(r"[\u0600-\u06ff]")
PDF_FONT_NAME: Final = "NotoSansArabic"


def safe_spreadsheet_value(value: str) -> str:
    """Prevent spreadsheet software from interpreting user text as a formula."""
    return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value


def _asset_path(relative: str) -> Path:
    return Path(settings.BASE_DIR) / "static" / relative


def render_xlsx(document: ReportDocument, *, language_code: str) -> bytes:
    """Render a bounded report to a styled, in-memory workbook."""
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = document.sheet_name or "Report"
    sheet.sheet_view.rightToLeft = language_code == "ar"
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.freeze_panes = "A6"

    logo = SpreadsheetImage(_asset_path("img/project-insight-logo.png"))
    logo.width = 190
    logo.height = 55
    sheet.add_image(logo, "A1")
    sheet.row_dimensions[1].height = 45

    column_count = max(len(document.headers), 1)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=column_count)
    title_cell = sheet.cell(2, 1, safe_spreadsheet_value(document.title))
    subtitle_cell = sheet.cell(3, 1, safe_spreadsheet_value(document.subtitle))
    title_cell.font = Font(color=PRIMARY, bold=True, size=18)
    subtitle_cell.font = Font(color=TEXT, italic=True, size=11)
    title_cell.alignment = Alignment(
        horizontal="right" if language_code == "ar" else "left"
    )
    subtitle_cell.alignment = Alignment(
        horizontal="right" if language_code == "ar" else "left"
    )

    header_row = 5
    for column, header in enumerate(document.headers, start=1):
        cell = sheet.cell(header_row, column, safe_spreadsheet_value(header))
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="right" if language_code == "ar" else "left",
            vertical="center",
        )
    sheet.row_dimensions[header_row].height = 24

    thin_sage = Side(style="thin", color=SAGE)
    for row_number, row in enumerate(document.rows, start=header_row + 1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, safe_spreadsheet_value(value))
            cell.alignment = Alignment(
                horizontal="right" if language_code == "ar" else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin_sage)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=SOFT_STONE)

    last_row = max(header_row, header_row + len(document.rows))
    last_column = get_column_letter(column_count)
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    assert sheet.oddFooter is not None
    assert sheet.oddFooter.center is not None
    assert sheet.oddFooter.right is not None
    sheet.oddFooter.center.text = "Project Insight"
    sheet.oddFooter.right.text = "Page &P of &N"

    for column, header in enumerate(document.headers, start=1):
        values = [header, *(row[column - 1] for row in document.rows)]
        width = min(max(max(len(str(value)) for value in values) + 2, 12), 38)
        sheet.column_dimensions[get_column_letter(column)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _register_pdf_font() -> None:
    if PDF_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(
                PDF_FONT_NAME,
                _asset_path("vendor/fonts/NotoSansArabic-VariableFont_wdth,wght.ttf"),
            )
        )


def _display_text(value: str, *, language_code: str) -> str:
    del language_code
    if ARABIC_PATTERN.search(value):
        return cast(str, get_display(arabic_reshaper.reshape(value)))
    return value


def _paragraph(
    value: str,
    *,
    language_code: str,
    style: ParagraphStyle,
) -> Paragraph:
    return Paragraph(
        escape(_display_text(value, language_code=language_code)),
        style,
    )


def render_pdf(document: ReportDocument, *, language_code: str) -> bytes:
    """Render a branded, embedded-font PDF entirely in memory."""
    _register_pdf_font()
    buffer = BytesIO()
    page_size = landscape(A4)
    report = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=15 * mm,
        title=document.title,
        author="Project Insight",
        subject=document.subtitle,
    )
    alignment: Literal[0, 2] = 2 if language_code == "ar" else 0
    normal = ParagraphStyle(
        "ReportNormal",
        fontName=PDF_FONT_NAME,
        fontSize=8,
        leading=11,
        alignment=alignment,
        textColor=colors.HexColor(f"#{TEXT}"),
        wordWrap="CJK",
    )
    heading = ParagraphStyle(
        "ReportHeading",
        parent=normal,
        fontSize=17,
        leading=21,
        textColor=colors.HexColor(f"#{PRIMARY}"),
    )
    subtitle = ParagraphStyle(
        "ReportSubtitle",
        parent=normal,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor(f"#{TEXT}"),
    )
    header = ParagraphStyle(
        "ReportHeader",
        parent=normal,
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )
    empty = ParagraphStyle(
        "ReportEmpty",
        parent=normal,
        fontSize=11,
        leading=15,
        textColor=colors.HexColor(f"#{TEXT}"),
    )

    logo = Image(
        str(_asset_path("img/project-insight-logo.png")),
        width=47 * mm,
        height=13.5 * mm,
    )
    title_block = [
        _paragraph(document.title, language_code=language_code, style=heading),
        _paragraph(document.subtitle, language_code=language_code, style=subtitle),
    ]
    if language_code == "ar":
        title_row = [logo, title_block]
        title_widths = [52 * mm, report.width - 52 * mm]
    else:
        title_row = [title_block, logo]
        title_widths = [report.width - 52 * mm, 52 * mm]
    title_table = Table([title_row], colWidths=title_widths)
    title_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "RIGHT" if language_code == "ar" else "LEFT",
                ),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story: list[Flowable] = [KeepTogether(title_table), Spacer(1, 4 * mm)]

    if document.rows:
        displayed_headers = (
            tuple(reversed(document.headers))
            if language_code == "ar"
            else document.headers
        )
        displayed_rows = (
            tuple(tuple(reversed(row)) for row in document.rows)
            if language_code == "ar"
            else document.rows
        )
        table_data = [
            [
                _paragraph(value, language_code=language_code, style=header)
                for value in displayed_headers
            ],
            *[
                [
                    _paragraph(value, language_code=language_code, style=normal)
                    for value in row
                ]
                for row in displayed_rows
            ],
        ]
        column_width = report.width / max(len(document.headers), 1)
        table = Table(
            table_data,
            colWidths=[column_width] * len(document.headers),
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{PRIMARY}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, -1), PDF_FONT_NAME),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "RIGHT" if language_code == "ar" else "LEFT",
                    ),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{SAGE}")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor(f"#{SOFT_STONE}")],
                    ),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(table)
    else:
        no_data = (
            "لا توجد بيانات مطابقة." if language_code == "ar" else "No matching data."
        )
        empty_table = Table(
            [[_paragraph(no_data, language_code=language_code, style=empty)]],
            colWidths=[report.width],
        )
        empty_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{SOFT_STONE}")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{STONE}")),
                    ("PADDING", (0, 0), (-1, -1), 12),
                ]
            )
        )
        story.append(empty_table)

    def add_footer(canvas: Any, doc: Any) -> None:
        del doc
        pdf_canvas = canvas
        pdf_canvas.saveState()
        pdf_canvas.setStrokeColor(colors.HexColor(f"#{SAGE}"))
        pdf_canvas.line(12 * mm, 10 * mm, page_size[0] - 12 * mm, 10 * mm)
        pdf_canvas.setFont(PDF_FONT_NAME, 7)
        pdf_canvas.setFillColor(colors.HexColor(f"#{TEXT}"))
        footer = f"Project Insight  |  {pdf_canvas.getPageNumber()}"
        pdf_canvas.drawCentredString(page_size[0] / 2, 6 * mm, footer)
        pdf_canvas.restoreState()

    report.build(story, onFirstPage=add_footer, onLaterPages=add_footer)
    return buffer.getvalue()
