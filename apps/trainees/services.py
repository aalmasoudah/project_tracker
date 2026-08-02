"""Transactional manual and import workflows for Phase 8."""

import csv
import hashlib
import io
import zipfile
from collections.abc import Iterable
from pathlib import Path

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import UploadedFile
from django.core.validators import validate_email
from django.db import transaction
from django.db.models import Max
from django.http import HttpRequest
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.audit import actions
from apps.audit.models import AuditEvent
from apps.audit.services import record_audit_event
from apps.courses.models import Course
from apps.trainees.models import CourseEnrollment, ImportBatch, ImportRow, Trainee
from apps.trainees.normalization import (
    normalize_trainee_phone,
    trainee_identity_key,
)
from apps.trainees.policies import (
    HEADER_ALIASES,
    IMPORT_MIME_TYPES,
    MAX_IMPORT_ROWS,
    MAX_IMPORT_SIZE,
    OPTIONAL_HEADERS,
    REQUIRED_HEADERS,
)
from apps.trainees.selectors import (
    can_archive_enrollment,
    can_manage_enrollment,
    ensure_import_scope,
)


def _audit(
    *,
    actor: User,
    action: str,
    target_type: str,
    target_id: str,
    target_label: str,
    metadata: dict[str, object] | None = None,
    request: HttpRequest | None = None,
) -> None:
    record_audit_event(
        actor=actor,
        action=action,
        target_type=target_type,
        target_id=target_id,
        target_label=target_label,
        metadata=metadata,
        request=request,
        scope=AuditEvent.Scope.TRAINEES,
    )


def _next_number(course: Course) -> int:
    highest = (
        CourseEnrollment.objects.filter(course=course).aggregate(
            value=Max("trainee_number")
        )["value"]
        or 0
    )
    return int(highest) + 1


def _validate_trainee(full_name: str, phone: str, email: str) -> None:
    if not full_name.strip():
        raise ValidationError(_("Full name is required."))
    if not normalize_trainee_phone(phone):
        raise ValidationError(_("Enter a valid phone number."))
    if email:
        validate_email(email)


@transaction.atomic
def create_enrollment(
    *,
    actor: User,
    course: Course,
    full_name: str,
    phone: str,
    email: str = "",
    request: HttpRequest | None = None,
) -> CourseEnrollment:
    course = (
        Course.objects.select_for_update().select_related("project").get(pk=course.pk)
    )
    if not can_manage_enrollment(actor, course):
        raise PermissionDenied(_("Trainee management permission is required."))
    _validate_trainee(full_name, phone, email)
    identity = trainee_identity_key(full_name, phone)
    if CourseEnrollment.objects.filter(course=course, identity_key=identity).exists():
        raise ValidationError(_("This trainee already exists in this course."))
    if (
        CourseEnrollment.objects.filter(course=course, is_archived=False).count()
        >= course.capacity
    ):
        raise ValidationError(_("Course capacity would be exceeded."))
    trainee = Trainee(
        full_name=full_name,
        phone=phone,
        email=email,
        created_by=actor,
        updated_by=actor,
    )
    trainee.full_clean()
    trainee.save()
    enrollment = CourseEnrollment(
        course=course,
        trainee=trainee,
        trainee_number=_next_number(course),
        created_by=actor,
    )
    enrollment.save()
    _audit(
        actor=actor,
        action=actions.TRAINEE_ENROLLED,
        target_type="enrollment",
        target_id=str(enrollment.pk),
        target_label=str(enrollment),
        request=request,
    )
    return enrollment


@transaction.atomic
def update_enrollment(
    *,
    actor: User,
    enrollment: CourseEnrollment,
    full_name: str,
    phone: str,
    email: str = "",
    request: HttpRequest | None = None,
) -> CourseEnrollment:
    enrollment = (
        CourseEnrollment.objects.select_for_update()
        .select_related("course__project", "trainee")
        .get(pk=enrollment.pk)
    )
    if enrollment.is_archived or not can_manage_enrollment(actor, enrollment.course):
        raise PermissionDenied(_("Trainee management permission is required."))
    _validate_trainee(full_name, phone, email)
    identity = trainee_identity_key(full_name, phone)
    if (
        CourseEnrollment.objects.filter(course=enrollment.course, identity_key=identity)
        .exclude(pk=enrollment.pk)
        .exists()
    ):
        raise ValidationError(_("This trainee already exists in this course."))
    trainee = enrollment.trainee
    trainee.full_name = full_name
    trainee.phone = phone
    trainee.email = email
    trainee.updated_by = actor
    trainee.full_clean()
    trainee.save()
    enrollment.identity_key = identity
    enrollment.save(update_fields=("identity_key",))
    _audit(
        actor=actor,
        action=actions.TRAINEE_UPDATED,
        target_type="enrollment",
        target_id=str(enrollment.pk),
        target_label=str(enrollment),
        request=request,
    )
    return enrollment


@transaction.atomic
def set_enrollment_archived(
    *,
    actor: User,
    enrollment: CourseEnrollment,
    archived: bool,
    request: HttpRequest | None = None,
) -> CourseEnrollment:
    enrollment = (
        CourseEnrollment.objects.select_for_update()
        .select_related("course__project", "trainee")
        .get(pk=enrollment.pk)
    )
    if not can_archive_enrollment(actor, enrollment):
        raise PermissionDenied(_("Trainee archive permission is required."))
    if enrollment.is_archived == archived:
        return enrollment
    if not archived:
        active_count = CourseEnrollment.objects.filter(
            course=enrollment.course, is_archived=False
        ).count()
        if active_count >= enrollment.course.capacity:
            raise ValidationError(_("Course capacity would be exceeded."))
    enrollment.is_archived = archived
    enrollment.archived_at = timezone.now() if archived else None
    enrollment.archived_by = actor if archived else None
    enrollment.save(update_fields=("is_archived", "archived_at", "archived_by"))
    _audit(
        actor=actor,
        action=actions.TRAINEE_ARCHIVED if archived else actions.TRAINEE_RESTORED,
        target_type="enrollment",
        target_id=str(enrollment.pk),
        target_label=str(enrollment),
        request=request,
    )
    return enrollment


def _canonical_headers(values: Iterable[object]) -> list[str]:
    headers: list[str] = []
    for value in values:
        raw = "" if value is None else str(value).strip().casefold()
        canonical = HEADER_ALIASES.get(raw, "")
        if not canonical:
            raise ValidationError(_("The import contains an unsupported header."))
        if canonical in headers:
            raise ValidationError(_("The import contains duplicate headers."))
        headers.append(canonical)
    if not REQUIRED_HEADERS.issubset(headers):
        raise ValidationError(_("The import is missing required headers."))
    if not set(headers).issubset(REQUIRED_HEADERS | OPTIONAL_HEADERS):
        raise ValidationError(_("The import contains an unsupported header."))
    return headers


def _read_csv(content: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValidationError(_("CSV files must use UTF-8 encoding.")) from error
    if "\x00" in text:
        raise ValidationError(_("The CSV file is malformed."))
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        raw_headers = next(reader)
    except StopIteration as error:
        raise ValidationError(_("The import file is empty.")) from error
    return _canonical_headers(raw_headers), [list(row) for row in reader]


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 1_000 or sum(item.file_size for item in members) > (
                25 * 1024 * 1024
            ):
                raise ValidationError(_("The Excel workbook is too large to process."))
            names = {name.casefold() for name in archive.namelist()}
            if any(
                name.endswith("vbaproject.bin")
                or "/externallinks/" in name
                or name.endswith(".bin")
                for name in names
            ):
                raise ValidationError(
                    _("Macros and external workbook links are not allowed.")
                )
    except zipfile.BadZipFile as error:
        raise ValidationError(_("The Excel file is malformed.")) from error
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as error:
        raise ValidationError(_("The Excel file is malformed.")) from error
    if len(workbook.sheetnames) != 1:
        workbook.close()
        raise ValidationError(_("The workbook must contain exactly one sheet."))
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as error:
        workbook.close()
        raise ValidationError(_("The import file is empty.")) from error
    data: list[list[object]] = []
    for row in rows:
        if any(isinstance(value, str) and value.startswith("=") for value in row):
            workbook.close()
            raise ValidationError(_("Spreadsheet formulas are not allowed."))
        data.append(list(row))
    workbook.close()
    return _canonical_headers(raw_headers), data


def validate_import_file(upload: UploadedFile) -> tuple[str, bytes]:
    filename = upload.name or ""
    size = upload.size or 0
    suffix = Path(filename).suffix.lower()
    if suffix not in IMPORT_MIME_TYPES:
        raise ValidationError(_("Only CSV and XLSX files are allowed."))
    if size <= 0 or size > MAX_IMPORT_SIZE:
        raise ValidationError(_("Import files must be 5 MB or smaller."))
    content_type = (upload.content_type or "").split(";")[0].strip().lower()
    if content_type and content_type not in IMPORT_MIME_TYPES[suffix]:
        raise ValidationError(_("The file content type does not match its extension."))
    content = upload.read(MAX_IMPORT_SIZE + 1)
    upload.seek(0)
    if len(content) != size or len(content) > MAX_IMPORT_SIZE:
        raise ValidationError(_("The uploaded file could not be read safely."))
    return suffix, content


@transaction.atomic
def preview_import(
    *,
    actor: User,
    course: Course,
    upload: UploadedFile,
    request: HttpRequest | None = None,
) -> ImportBatch:
    course = Course.objects.select_related("project").get(pk=course.pk)
    ensure_import_scope(actor, course)
    suffix, content = validate_import_file(upload)
    headers, rows = _read_csv(content) if suffix == ".csv" else _read_xlsx(content)
    indexed_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows, start=2)
        if any(str(value).strip() for value in row if value is not None)
    ]
    if not indexed_rows:
        raise ValidationError(_("The import contains no trainee rows."))
    if len(indexed_rows) > MAX_IMPORT_ROWS:
        raise ValidationError(_("The import may contain at most 5000 rows."))
    batch = ImportBatch.objects.create(
        course=course,
        original_name=Path(upload.name or "").name[:255],
        file_type=suffix.lstrip("."),
        size=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        uploaded_by=actor,
    )
    existing = {
        enrollment.identity_key: enrollment
        for enrollment in CourseEnrollment.objects.filter(course=course).select_related(
            "trainee"
        )
    }
    seen: dict[str, CourseEnrollment | None] = {}
    valid_count = warning_count = duplicate_count = error_count = 0
    import_rows: list[ImportRow] = []
    for row_number, raw_values in indexed_rows:
        values = {
            header: (
                ""
                if index >= len(raw_values) or raw_values[index] is None
                else str(raw_values[index]).strip()
            )
            for index, header in enumerate(headers)
        }
        full_name = values.get("full_name", "")
        phone = values.get("phone", "")
        email = values.get("email", "")
        identity = trainee_identity_key(full_name, phone)
        classification = ImportRow.Classification.VALID
        resolution: str = ImportRow.Resolution.CREATE
        error_message = ""
        duplicate = existing.get(identity)
        try:
            if len(raw_values) != len(headers):
                raise ValidationError(
                    _("The row does not match the number of headers.")
                )
            _validate_trainee(full_name, phone, email)
        except ValidationError as error:
            classification = ImportRow.Classification.ERROR
            resolution = ""
            error_message = " ".join(error.messages)
            error_count += 1
        else:
            if duplicate is not None or identity in seen:
                classification = ImportRow.Classification.DUPLICATE
                resolution = ImportRow.Resolution.SKIP
                duplicate_count += 1
            elif not email:
                classification = ImportRow.Classification.WARNING
                warning_count += 1
            else:
                valid_count += 1
            seen[identity] = duplicate
        import_rows.append(
            ImportRow(
                batch=batch,
                row_number=row_number,
                full_name=full_name,
                phone=phone,
                email=email,
                identity_key=identity,
                classification=classification,
                resolution=resolution,
                error_message=error_message,
                duplicate_enrollment=duplicate,
            )
        )
    ImportRow.objects.bulk_create(import_rows)
    batch.row_count = len(import_rows)
    batch.valid_count = valid_count
    batch.warning_count = warning_count
    batch.duplicate_count = duplicate_count
    batch.error_count = error_count
    batch.save(
        update_fields=(
            "row_count",
            "valid_count",
            "warning_count",
            "duplicate_count",
            "error_count",
        )
    )
    _audit(
        actor=actor,
        action=actions.TRAINEE_IMPORT_PREVIEWED,
        target_type="import_batch",
        target_id=str(batch.pk),
        target_label=batch.original_name,
        metadata={"rows": batch.row_count, "errors": batch.error_count},
        request=request,
    )
    return batch


@transaction.atomic
def confirm_import(
    *,
    actor: User,
    batch: ImportBatch,
    resolutions: dict[int, str] | None = None,
    request: HttpRequest | None = None,
) -> ImportBatch:
    batch = (
        ImportBatch.objects.select_for_update()
        .select_related("course__project")
        .get(pk=batch.pk)
    )
    course = Course.objects.select_for_update().get(pk=batch.course_id)
    batch.course = course
    ensure_import_scope(actor, course, confirm=True)
    if batch.status != ImportBatch.Status.PREVIEW:
        raise ValidationError(_("This import is no longer awaiting confirmation."))
    if batch.error_count:
        raise ValidationError(_("Resolve all import errors before confirmation."))
    rows = list(batch.rows.select_for_update().order_by("row_number"))
    resolutions = resolutions or {}
    create_rows: list[ImportRow] = []
    for row in rows:
        requested = resolutions.get(row.pk, row.resolution)
        if row.classification == ImportRow.Classification.DUPLICATE:
            if requested not in (
                ImportRow.Resolution.SKIP,
                ImportRow.Resolution.UPDATE,
            ):
                raise ValidationError(_("Choose Skip or Update for every duplicate."))
            row.resolution = requested
            if (
                requested == ImportRow.Resolution.UPDATE
                and row.duplicate_enrollment_id is None
            ):
                raise ValidationError(
                    _("Duplicates inside the same file can only be skipped.")
                )
        elif row.classification in (
            ImportRow.Classification.VALID,
            ImportRow.Classification.WARNING,
        ):
            row.resolution = ImportRow.Resolution.CREATE
            create_rows.append(row)
    active_count = CourseEnrollment.objects.filter(
        course=course, is_archived=False
    ).count()
    if active_count + len(create_rows) > course.capacity:
        raise ValidationError(_("Course capacity would be exceeded."))
    next_number = _next_number(course)
    for row in rows:
        if row.resolution == ImportRow.Resolution.SKIP:
            row.save(update_fields=("resolution",))
            continue
        if row.resolution == ImportRow.Resolution.UPDATE:
            duplicate = row.duplicate_enrollment
            if duplicate is None:
                raise ValidationError(_("The duplicate record changed after preview."))
            duplicate = (
                CourseEnrollment.objects.select_for_update()
                .select_related("trainee")
                .get(pk=duplicate.pk, course=course, identity_key=row.identity_key)
            )
            trainee = duplicate.trainee
            trainee.full_name = row.full_name
            trainee.phone = row.phone
            trainee.email = row.email
            trainee.updated_by = actor
            trainee.full_clean()
            trainee.save()
            row.save(update_fields=("resolution",))
            continue
        if CourseEnrollment.objects.filter(
            course=course, identity_key=row.identity_key
        ).exists():
            raise ValidationError(
                _("Trainee data changed after preview; preview again.")
            )
        trainee = Trainee(
            full_name=row.full_name,
            phone=row.phone,
            email=row.email,
            created_by=actor,
            updated_by=actor,
        )
        trainee.full_clean()
        trainee.save()
        enrollment = CourseEnrollment.objects.create(
            course=course,
            trainee=trainee,
            trainee_number=next_number,
            created_by=actor,
        )
        next_number += 1
        row.created_enrollment = enrollment
        row.save(update_fields=("resolution", "created_enrollment"))
    batch.status = ImportBatch.Status.CONFIRMED
    batch.confirmed_by = actor
    batch.confirmed_at = timezone.now()
    batch.save(update_fields=("status", "confirmed_by", "confirmed_at"))
    _audit(
        actor=actor,
        action=actions.TRAINEE_IMPORT_CONFIRMED,
        target_type="import_batch",
        target_id=str(batch.pk),
        target_label=batch.original_name,
        metadata={"rows": batch.row_count},
        request=request,
    )
    return batch


@transaction.atomic
def cancel_import(
    *,
    actor: User,
    batch: ImportBatch,
    request: HttpRequest | None = None,
) -> ImportBatch:
    batch = (
        ImportBatch.objects.select_for_update()
        .select_related("course__project")
        .get(pk=batch.pk)
    )
    ensure_import_scope(actor, batch.course)
    if batch.status != ImportBatch.Status.PREVIEW:
        raise ValidationError(_("This import is no longer awaiting confirmation."))
    batch.status = ImportBatch.Status.CANCELLED
    batch.save(update_fields=("status",))
    _audit(
        actor=actor,
        action=actions.TRAINEE_IMPORT_CANCELLED,
        target_type="import_batch",
        target_id=str(batch.pk),
        target_label=batch.original_name,
        request=request,
    )
    return batch
